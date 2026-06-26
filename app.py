"""
Generador de Informes - Distrito 17D08 v3
"""
import os, json, copy, webbrowser, threading
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, session
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

app = Flask(__name__)
app.secret_key = 'informe-17d08-local'
BASE = os.path.dirname(os.path.abspath(__file__))
app.config['UPLOAD_FOLDER'] = os.path.join(BASE, 'uploads')
app.config['OUTPUT_FOLDER'] = os.path.join(BASE, 'output')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024
PLANTILLA_PATH = os.path.join(BASE, 'Plantilla.docx')
for d in [app.config['UPLOAD_FOLDER'], app.config['OUTPUT_FOLDER']]:
    os.makedirs(d, exist_ok=True)

# ─── UTILS ────────────────────────────────────────────────────────────────────
def clean_id(val):
    s = str(val).strip()
    if s.endswith('.0'): s = s[:-2]
    if s.isdigit() and len(s) < 10: s = s.zfill(10)
    return s

def read_distributivo(fp):
    try: df = pd.read_excel(fp, dtype={'NÚMERO IDENTIFICACIÓN': str, 'ESTRUCTURA PROGRAMÁTICA': str})
    except: df = pd.read_excel(fp)
    if 'NÚMERO IDENTIFICACIÓN' in df.columns:
        df['NÚMERO IDENTIFICACIÓN'] = df['NÚMERO IDENTIFICACIÓN'].apply(clean_id)
    if 'CÓDIGO ENLACE PRESUPUESTARIO' in df.columns:
        df['CÓDIGO ENLACE PRESUPUESTARIO'] = pd.to_numeric(df['CÓDIGO ENLACE PRESUPUESTARIO'], errors='coerce').fillna(0).astype(int)
    return df

def filter_ocupados(df):
    return df[df['ESTADO DEL PUESTO'].str.strip().str.upper() == 'OCUPADO'].copy()

def compare_distributivos(df_ant, df_act):
    ant, act = filter_ocupados(df_ant), filter_ocupados(df_act)
    ids_ant, ids_act = set(ant['NÚMERO IDENTIFICACIÓN'].unique()), set(act['NÚMERO IDENTIFICACIÓN'].unique())
    sal = ant[ant['NÚMERO IDENTIFICACIÓN'].isin(ids_ant - ids_act)].copy(); sal['MOTIVO'] = 'RENUNCIA'
    ent = act[act['NÚMERO IDENTIFICACIÓN'].isin(ids_act - ids_ant)].copy(); ent['MOTIVO'] = 'VINCULACIÓN'
    # Detect people present in both but with RMU or modalidad change (won concurso)
    comunes = ids_ant & ids_act
    ant_idx = ant[ant['NÚMERO IDENTIFICACIÓN'].isin(comunes)].set_index('NÚMERO IDENTIFICACIÓN')
    act_idx = act[act['NÚMERO IDENTIFICACIÓN'].isin(comunes)].set_index('NÚMERO IDENTIFICACIÓN')
    cambio_ids = []
    for ced in comunes:
        try:
            rmu_a = float(ant_idx.loc[ced, 'RMU PUESTO'])
            rmu_b = float(act_idx.loc[ced, 'RMU PUESTO'])
            mod_a = str(ant_idx.loc[ced, 'MODALIDAD LABORAL']).strip().upper()
            mod_b = str(act_idx.loc[ced, 'MODALIDAD LABORAL']).strip().upper()
            if rmu_a != rmu_b or mod_a != mod_b:
                cambio_ids.append(ced)
        except: pass
    camb = act[act['NÚMERO IDENTIFICACIÓN'].isin(cambio_ids)].copy()
    camb['MOTIVO'] = 'VINCULACIÓN GANADOR DE CONCURSO'
    return sal, ent, camb

def get_distribution(df_act):
    ocu = filter_ocupados(df_act)
    n = lambda s: str(s).strip().upper()
    fr = lambda df, kw: df[df['RÉGIMEN LABORAL'].apply(n).str.contains(kw)]
    fm = lambda df, kw: df[df['MODALIDAD LABORAL'].apply(n).str.contains(kw)]
    losep, otros, ct = fr(ocu,'LOSEP'), fr(ocu,'OTRO'), fr(ocu,'CODIGO DEL TRABAJO')
    oc = fm(otros,'CONTRATO')
    cats = {
        '9.1':{'title':'PERSONAL ADMINISTRATIVO - CONTRATOS OCASIONALES LOSEP','subtitle':'a) PROGRAMA: 51','data':fm(losep,'CONTRATO')},
        '9.2':{'title':'NÓMINA ADMINISTRATIVOS BAJO LA MODALIDAD NOMBRAMIENTO','subtitle':'b) PROGRAMA 01 GRUPO 51','data':fm(losep,'NOMBRAMIENTO')},
        '9.3':{'title':'NÓMINA DE DOCENTES BAJO LA MODALIDAD DE CONTRATO OCASIONAL','subtitle':'c) PROGRAMA 55-56-57 GRUPO 51','data':oc[oc['CÓDIGO ENLACE PRESUPUESTARIO'].isin([13,14,15,19])]},
        '9.4':{'title':'NÓMINA DE DOCENTES BAJO LA MODALIDAD DE CONTRATO OCASIONAL - EBJAS','subtitle':'d) PROGRAMA 55-58 GRUPO 51 EBJAS','data':oc[oc['CÓDIGO ENLACE PRESUPUESTARIO'].isin([6,7,8])]},
        '9.5':{'title':'NÓMINA DE DOCENTES BAJO NOMBRAMIENTO','subtitle':'e) PROGRAMA 56-57 GRUPO 51','data':fm(otros,'NOMBRAMIENTO')},
        '9.6':{'title':'NÓMINA AL CÓDIGO DE TRABAJO','subtitle':'f) PROGRAMA 01 - GRUPO DE GASTO 51','data':ct.copy()},
    }
    for k in cats: cats[k]['data'] = cats[k]['data'].sort_values('NOMBRES').reset_index(drop=True)
    return cats

def fmt_rmu(v):
    try: return f"{float(v):,.2f}"
    except: return str(v)

def df_to_rows(df):
    return [{'nro':i+1,'cedula':clean_id(r.get('NÚMERO IDENTIFICACIÓN','')),'nombres':str(r.get('NOMBRES','')).strip(),
             'escala':str(r.get('ESCALA OCUPACIONAL','')).strip(),'rmu':fmt_rmu(r.get('RMU PUESTO','')),'observacion':''}
            for i,(_,r) in enumerate(df.iterrows())]

# ─── WORD HELPERS ─────────────────────────────────────────────────────────────
def sc(cell, color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>'))

def set_col_width(cell, width_cm):
    """Set explicit column width on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(width_cm * 567)}" w:type="dxa"/>')
        tcPr.append(tcW)
    else:
        tcW.set(qn('w:w'), str(int(width_cm * 567)))
        tcW.set(qn('w:type'), 'dxa')

def set_grid(table, widths_cm):
    """Define los anchos de columna UNA sola vez vía tblGrid (rápido).
    Reemplaza el anti-patrón set_col_width por celda (que en tablas grandes
    hacía miles de operaciones XML y tardaba minutos)."""
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is None:
        return
    for gc in list(grid.findall(qn('w:gridCol'))):
        grid.remove(gc)
    for w in widths_cm:
        grid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{int(w * 567)}"/>'))

def _esc_xml(s):
    return (str(s) if s is not None else '').replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('"','&quot;')

def _tr_xml(vals, widths_dxa, sz=8, align='center'):
    """XML de una fila de datos (w:tr) construido directo. Aptos, centrado, bordes via estilo de tabla."""
    SZ = str(int(sz * 2))
    cells = []
    for i, v in enumerate(vals):
        w = widths_dxa[i] if widths_dxa and i < len(widths_dxa) else None
        tcPr = '<w:tcPr>' + (f'<w:tcW w:w="{w}" w:type="dxa"/>' if w else '') + '<w:vAlign w:val="center"/></w:tcPr>'
        rpr = f'<w:rPr><w:rFonts w:ascii="Aptos" w:hAnsi="Aptos"/><w:sz w:val="{SZ}"/></w:rPr>'
        ppr = f'<w:pPr><w:jc w:val="{align}"/><w:spacing w:before="20" w:after="20"/></w:pPr>'
        cells.append(f'<w:tc>{tcPr}<w:p>{ppr}<w:r>{rpr}<w:t xml:space="preserve">{_esc_xml(v)}</w:t></w:r></w:p></w:tc>')
    return f'<w:tr>{"".join(cells)}</w:tr>'

def append_rows_fast(table, rows_vals, widths_dxa, sz=8):
    """Inyecta muchas filas de datos vía XML directo en una sola operación.
    Evita el overhead de python-docx celda-por-celda (que en tablas grandes
    tardaba minutos). Los bordes los hereda del estilo 'Table Grid' de la tabla."""
    xml = ''.join(_tr_xml(v, widths_dxa, sz=sz) for v in rows_vals)
    frag = parse_xml(f'<w:root {nsdecls("w")}>{xml}</w:root>')
    tbl = table._tbl
    for tr in list(frag):
        tbl.append(tr)

def ct(cell, text, bold=False, sz=8, al=WD_ALIGN_PARAGRAPH.CENTER, fn='Aptos'):
    cell.text = ''
    p = cell.paragraphs[0]; p.alignment = al
    pf = p.paragraph_format; pf.space_before = Pt(1); pf.space_after = Pt(1)
    r = p.add_run(str(text)); r.font.size = Pt(sz); r.font.name = fn; r.font.bold = bold

def ap(doc, text, bold=False, sz=10, al=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False):
    p = doc.add_paragraph(); p.alignment = al
    pf = p.paragraph_format; pf.space_before = Pt(2); pf.space_after = Pt(2)
    r = p.add_run(text); r.font.size = Pt(sz); r.font.name = 'Aptos'; r.font.bold = bold; r.font.italic = italic
    return p

def set_table_width(table, width_cm):
    """Set total table width."""
    tblPr = table._tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}/>')
        table._tbl.insert(0, tblPr)
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="{int(width_cm * 567)}" w:type="dxa"/>')
        tblPr.append(tblW)
    else:
        tblW.set(qn('w:w'), str(int(width_cm * 567)))
        tblW.set(qn('w:type'), 'dxa')

def add_compact_table(doc, rows, tnum, ttitle, sub=''):
    """Standard 6-col table matching example informe width (~13.5cm centered on page)."""
    ap(doc, f'{tnum}.- {ttitle}', bold=True, sz=10)
    if sub: ap(doc, sub, bold=True, sz=9)
    if not rows:
        ap(doc, 'Sin novedades en este período.', sz=9, italic=True); return

    hdrs = ['NRO.','CÉDULA','APELLIDOS Y NOMBRES','ESCALA','RMU','OBSERVACIÓN']
    # Widths in cm matching example: total ~13.5cm
    widths = [1.0, 1.6, 3.5, 3.0, 1.4, 3.0]
    widths_dxa = [int(w * 567) for w in widths]
    t = doc.add_table(rows=1, cols=6)   # solo el header; las filas de datos se inyectan por XML
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    set_table_width(t, sum(widths))
    set_grid(t, widths)   # anchos definidos UNA vez (rápido)

    for i, h in enumerate(hdrs):
        c = t.rows[0].cells[i]
        ct(c, h, bold=True, sz=8)
        sc(c, 'D9D9D9')   # cabecera pintada (gris, estilo Fondos)

    # Filas de datos vía XML directo (13x más rápido que celda-por-celda)
    rows_vals = [[r['nro'], r['cedula'], r['nombres'], r['escala'], r['rmu'], r.get('observacion', '')]
                 for r in rows]
    append_rows_fast(t, rows_vals, widths_dxa, sz=8)

    doc.add_paragraph()

def add_manual_nov(doc, nov, num):
    ap(doc, f'{num}.- NOVEDAD: {nov.get("title","")}', bold=True, sz=10)
    if nov.get('description'): ap(doc, nov['description'], sz=9)
    rows = nov.get('rows', []); cols = nov.get('columns', [])
    if not rows:
        ap(doc, 'Sin novedades en este período.', sz=9, italic=True); return
    nc = len(cols)
    t = doc.add_table(rows=1, cols=nc)   # solo header; datos por XML
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    # ~13.5cm total, distributed evenly
    total_w = 13.5
    col_w = total_w / nc
    widths_dxa = [int(col_w * 567)] * nc
    set_table_width(t, total_w)
    set_grid(t, [col_w] * nc)   # anchos una vez (rápido)
    for i, h in enumerate(cols):
        ct(t.rows[0].cells[i], h, bold=True, sz=8)
        sc(t.rows[0].cells[i], 'D9D9D9')   # cabecera pintada
    # Filas de datos vía XML directo
    rows_vals = [[(row[ci] if ci < len(row) else '') for ci in range(nc)] for row in rows]
    append_rows_fast(t, rows_vals, widths_dxa, sz=8)
    doc.add_paragraph()

def gen_datos_table(doc, d):
    t = doc.add_table(rows=8, cols=6); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    set_table_width(t, 16.0)
    t.rows[0].cells[0].merge(t.rows[0].cells[5])
    ct(t.rows[0].cells[0],'DATOS GENERALES',bold=True,sz=8,al=WD_ALIGN_PARAGRAPH.CENTER); sc(t.rows[0].cells[0],'BEBEBE')
    ct(t.rows[1].cells[0],'Fecha de Informe',bold=True,sz=8); sc(t.rows[1].cells[0],'D9D9D9')
    t.rows[1].cells[1].merge(t.rows[1].cells[3]); ct(t.rows[1].cells[1],d.get('fecha',''),sz=8)
    t.rows[1].cells[4].merge(t.rows[1].cells[5]); ct(t.rows[1].cells[4],d.get('numero_informe',''),bold=True,sz=8)
    for ri in [2,5]:
        lbl = 'Funcionario Responsable de Informe' if ri==2 else 'Informe dirigido a:'
        ct(t.rows[ri].cells[0],lbl,bold=True,sz=7); sc(t.rows[ri].cells[0],'D9D9D9')
        for h,ci in [('Nombre',1),('Extensión',2),('Correo',3),('Cargo',5)]:
            ct(t.rows[ri].cells[ci],h,bold=True,sz=7); sc(t.rows[ri].cells[ci],'D9D9D9')
    for prefix,row_i in [('responsable',4),('dirigido',7)]:
        sc(t.rows[row_i].cells[0],'D9D9D9')
        ct(t.rows[row_i].cells[1],d.get(f'{prefix}_nombre',''),sz=8)
        ct(t.rows[row_i].cells[2],d.get(f'{prefix}_ext',''),sz=8)
        ct(t.rows[row_i].cells[3],d.get(f'{prefix}_email',''),sz=8)
        ct(t.rows[row_i].cells[5],d.get(f'{prefix}_cargo',''),sz=8)
    doc.add_paragraph()

def gen_firma_table(doc, d):
    t = doc.add_table(rows=6, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    set_table_width(t, 13.5)
    fw = [5.0, 4.0, 4.5]
    set_grid(t, fw)
    for bs, nk, ck in [(0,'elaborado_nombre','elaborado_cargo'),(3,'aprobado_nombre','aprobado_cargo')]:
        t.rows[bs].cells[0].merge(t.rows[bs].cells[2])
        ct(t.rows[bs].cells[0],'DESARROLLO DEL DOCUMENTO',sz=7,al=WD_ALIGN_PARAGRAPH.CENTER); sc(t.rows[bs].cells[0],'BEBEBE')
        for i,lbl in enumerate(['Nombre','Firma','Fecha']):
            ct(t.rows[bs+1].cells[i],lbl,sz=7,al=WD_ALIGN_PARAGRAPH.CENTER); sc(t.rows[bs+1].cells[i],'D9D9D9')
        cell = t.rows[bs+2].cells[0]; cell.text = ''
        p1 = cell.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(d.get(nk,'')); r1.font.size = Pt(7); r1.font.name = 'Aptos'
        p2 = cell.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(d.get(ck,'')); r2.font.size = Pt(7); r2.font.name = 'Aptos'
        ct(t.rows[bs+2].cells[1],'',sz=7)
        ct(t.rows[bs+2].cells[2],d.get('fecha',''),sz=7,al=WD_ALIGN_PARAGRAPH.CENTER)

def gen_conclusion_table(doc, cats):
    c = {k:len(cats[k]['data']) for k in cats}
    tl = c['9.1']+c['9.2']; td = c['9.3']+c['9.4']+c['9.5']; total = tl+c['9.6']+td
    t = doc.add_table(rows=7, cols=7); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    set_table_width(t, 13.5)
    t.rows[0].cells[0].merge(t.rows[0].cells[6])
    ct(t.rows[0].cells[0],'PROGRAMAS',bold=True,sz=8); sc(t.rows[0].cells[0],'D9D9D9')
    for i,h in enumerate(['PERSONAL','LOSEP','CÓD. TRABAJO','56-001 G51','PROG 55-57/G51','PROG 58/55','TOTAL']):
        ct(t.rows[1].cells[i],h,bold=True,sz=7)
    for ri,rd in enumerate([['ADM. NOMBRAMIENTOS',str(c['9.2']),'','','','0',str(c['9.2'])],['ADM. CONTRATO OCASIONAL',str(c['9.1']),'','','','',str(c['9.1'])],['CÓDIGO DE TRABAJO','',str(c['9.6']),'','0','0',str(c['9.6'])],['DOCENTES','','',str(c['9.5']),str(c['9.3']),str(c['9.4']),str(td)]]):
        for ci,v in enumerate(rd):
            ct(t.rows[ri+2].cells[ci],v,sz=8)
    for ci,v in enumerate(['TOTAL',str(tl),str(c['9.6']),str(c['9.5']),str(c['9.3']),str(c['9.4']),str(total)]):
        ct(t.rows[6].cells[ci],v,bold=True,sz=8)
    return total

def generate_word(dg, novs, cats, out):
    if os.path.exists(PLANTILLA_PATH):
        doc = Document(PLANTILLA_PATH)
        for el in list(doc.element.body):
            tag = el.tag.split('}')[-1] if '}' in el.tag else el.tag
            if tag != 'sectPr': doc.element.body.remove(el)
    else: doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Aptos'; st.font.size = Pt(10)
    mes, anio = dg.get('mes','MES').upper(), dg.get('anio',datetime.now().year)
    gen_datos_table(doc, dg)
    ap(doc,'OBJETIVO:',bold=True,sz=10)
    ap(doc,f'Emitir el informe técnico de novedades de nómina presentadas en el MES DE {mes} DE {anio}.',sz=10)
    ap(doc,'BASE LEGAL',bold=True,sz=10)
    for b,txt in [(False,'Acuerdo Interministerial No. MDT-2015-0002 entre el Ministerio de Trabajo y el Ministerio de Finanzas- Subrogante en su disposición tercera.'),(True,'LOSEP:'),(False,'Art. 106.- Pago por remuneraciones - El pago de remuneraciones se hará por mensualidades o quincenas.'),(True,'REGLAMENTO DE LA LOSEP:'),(False,'Art. 254.- Pago de remuneraciones. - El pago de las remuneraciones mensuales unificadas de las y los servidores públicos que laboren bajo la modalidad de contrato de servicios ocasionales, se efectuará en forma mensual o quincenal conforme a la normativa vigente.'),(False,'La remuneración de la o el servidor público, que estuviere en el ejercicio de un puesto, será pagada desde el primer día del mes, y en el caso de cesación se pagará hasta el último día del mes correspondiente.')]:
        ap(doc,txt,bold=b,sz=10 if b else 9)
    ap(doc,'DESARROLLO:',bold=True,sz=11); ap(doc,'NOVEDADES:',bold=True,sz=10)
    ic = 1
    for nv in [n for n in novs if n.get('active')]:
        add_manual_nov(doc, nv, ic); ic += 1
    dn = ic
    ap(doc,f'{dn}.- NOVEDADES NÓMINA DE FUNCIONARIOS PERTENECIENTES A LA DIRECCIÓN DISTRITAL:',bold=True,sz=10)
    for i,k in enumerate(['9.1','9.2','9.3','9.4','9.5','9.6']):
        add_compact_table(doc, df_to_rows(cats[k]['data']), f'{dn}.{i+1}', cats[k]['title'], cats[k]['subtitle'])
    ic = dn + 1
    ap(doc,f'{ic}.- CONCLUSIONES. -',bold=True,sz=10)
    total = gen_conclusion_table(doc, cats)
    ap(doc,f'En base a lo anteriormente citado la Unidad Distrital de Talento Humano que la División Distrital Administrativa Financiera debe realizar el trámite respectivo para la cancelación de remuneraciones del mes de {mes} DE {anio} del personal bajo el régimen LOSEP, LOEI y Código de Trabajo, según el detalle presentado.',sz=10)
    ic += 1; ap(doc,f'{ic}.- RECOMENDACIONES. -',bold=True,sz=10)
    ap(doc,f'La Unidad Distrital de Talento Humano recomienda a la División Distrital Administrativa Financiera realice el trámite pertinente para el pago de haberes del mes de {mes} {anio} a los {total} funcionarios que laboran en la Dirección Distrital 17D08 Parroquias Rurales Conocoto a la Merced.',sz=10)
    doc.add_paragraph(); doc.add_paragraph()
    gen_firma_table(doc, dg); doc.save(out)

# ─── EXCEL ────────────────────────────────────────────────────────────────────
def generate_excel(cats, novs, out):
    wb = Workbook(); hf = Font(name='Arial Narrow',bold=True,size=9); df_ = Font(name='Arial Narrow',size=9)
    bd = Border(left=Side('thin'),right=Side('thin'),top=Side('thin'),bottom=Side('thin'))
    def wt(ws, rows, sr=1):
        for c,h in enumerate(['NRO.','CÉDULA','APELLIDOS Y NOMBRES','ESCALA','RMU','OBSERVACIÓN'],1):
            cl=ws.cell(row=sr,column=c,value=h); cl.font=hf; cl.border=bd; cl.alignment=Alignment(horizontal='center')
        for i,row in enumerate(rows):
            for c,v in enumerate([row['nro'],row['cedula'],row['nombres'],row['escala'],row['rmu'],row.get('observacion','')],1):
                cl=ws.cell(row=sr+1+i,column=c,value=v); cl.font=df_; cl.border=bd
        ws.column_dimensions['A'].width=6; ws.column_dimensions['B'].width=15; ws.column_dimensions['C'].width=40
        ws.column_dimensions['D'].width=35; ws.column_dimensions['E'].width=12; ws.column_dimensions['F'].width=30
    wb.remove(wb.active)
    for nv in [n for n in novs if n.get('active') and n.get('rows')]:
        sn=nv['title'][:28].replace('/','_'); ws=wb.create_sheet(title=sn)
        for c,h in enumerate(nv.get('columns',[]),1): cl=ws.cell(row=1,column=c,value=h); cl.font=hf; cl.border=bd
        for i,row in enumerate(nv['rows']):
            for c,v in enumerate(row,1): cl=ws.cell(row=2+i,column=c,value=v); cl.font=df_; cl.border=bd
    nm={'9.1':'Adm Contrato LOSEP','9.2':'Adm Nombr LOSEP','9.3':'Doc Contrato 13-19','9.4':'Doc Contrato 6-8','9.5':'Doc Nombramiento','9.6':'Código Trabajo'}
    for k in ['9.1','9.2','9.3','9.4','9.5','9.6']:
        ws=wb.create_sheet(title=nm[k]); wt(ws,df_to_rows(cats[k]['data']))
    ws=wb.create_sheet(title='RESUMEN')
    ws.cell(row=1,column=1,value='CATEGORÍA').font=hf; ws.cell(row=1,column=2,value='TOTAL').font=hf
    ws.column_dimensions['A'].width=40; ws.column_dimensions['B'].width=12
    for i,k in enumerate(['9.1','9.2','9.3','9.4','9.5','9.6']):
        ws.cell(row=2+i,column=1,value=cats[k]['title']).font=df_; ws.cell(row=2+i,column=2,value=len(cats[k]['data'])).font=df_
    t=sum(len(cats[k]['data']) for k in cats)
    ws.cell(row=9,column=1,value='TOTAL GENERAL').font=Font(name='Arial Narrow',bold=True,size=10)
    ws.cell(row=9,column=2,value=t).font=Font(name='Arial Narrow',bold=True,size=10)
    wb.save(out)

# ─── PREVIOUS DATA ────────────────────────────────────────────────────────────
# Las novedades precargadas viven en novedades_anteriores.json (NO versionado:
# contiene cedulas/nombres reales). Copia novedades_anteriores.example.json a
# novedades_anteriores.json para usar el atajo. Si no existe, queda vacio y la
# app funciona igual (sin precarga).
PREV_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'novedades_anteriores.json')

def _load_prev():
    try:
        if os.path.exists(PREV_PATH):
            with open(PREV_PATH, encoding='utf-8') as f:
                return json.load(f)
    except Exception as e:
        print(f"[AVISO] No se pudo leer novedades_anteriores.json: {e}")
    return {}

DEFAULT_NOVEDADES = [
    {'id':'encargos','title':'ENCARGOS DE FUNCIONES DE DIRECTORES CON PARTIDA DE DOCENTE Y LA DIFERENCIA CON LA PARTIDA DE ENCARGO','description':'Detalle de docentes con encargo de funciones de director/a.','active':False,'columns':['No','CÉDULA','APELLIDOS Y NOMBRES','INSTITUCIÓN','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
    {'id':'jubilacion','title':'JUBILACIÓN VOLUNTARIA','description':'','active':False,'columns':['No','APELLIDOS Y NOMBRES','C.C. No.','CARGO','RMU','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
    {'id':'cambio_regimen','title':'CAMBIO DE RÉGIMEN LABORAL DE LOSEP A LOEI','description':'','active':False,'columns':['NRO.','CÉDULA','APELLIDOS Y NOMBRES','ESCALA','RMU','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
    {'id':'incorporacion_contrato','title':'INCORPORACIÓN POR EDUCA-EMPLEO CONTRATO','description':'','active':False,'columns':['NRO','CÉDULA','APELLIDOS Y NOMBRES','ESCALA','RMU','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
    {'id':'incorporacion_np','title':'INCORPORACIÓN NOMBRAMIENTO PROVISIONAL','description':'','active':False,'columns':['NRO.','APELLIDOS Y NOMBRES','CÉDULA','ESCALA','FECHA INGRESO','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
    {'id':'suspension','title':'SUSPENSIÓN SIN REMUNERACIÓN','description':'','active':False,'columns':['NRO.','CÉDULA','NOMBRES','ESCALA','RMU','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
    {'id':'fallecimiento','title':'FALLECIMIENTO','description':'','active':False,'columns':['NRO.','CÉDULA','NOMBRES','ESCALA','RMU','OBSERVACIÓN'],'rows':[],'has_previous':False,'is_auto':True},
    {'id':'renuncia','title':'RENUNCIA','description':'','active':False,'columns':['NRO.','CÉDULA','NOMBRES','ESCALA','RMU','OBSERVACIÓN'],'rows':[],'has_previous':False,'is_auto':True},
    {'id':'notificacion','title':'NOTIFICACIÓN','description':'','active':False,'columns':['NRO.','CÉDULA','NOMBRES','ESCALA','RMU','OBSERVACIÓN'],'rows':[],'has_previous':False,'is_auto':True},
    {'id':'reforma_aprobada','title':'FUNCIONARIA REFORMA APROBADA','description':'','active':False,'columns':['NRO.','CÉDULA','APELLIDOS Y NOMBRES','ESCALA','RMU','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
    {'id':'jubilados_pension','title':'LISTADO DE JUBILADOS – PENSIÓN JUBILAR','description':'','active':False,'columns':['No.','CÉDULA','APELLIDOS Y NOMBRES','VALIDACIÓN SUPERVIVENCIA','OBSERVACIÓN'],'rows':[],'has_previous':True,'is_auto':False},
]

# ─── ROUTES ───────────────────────────────────────────────────────────────────
@app.route('/')
def index(): return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload():
    fa, fb = request.files.get('distributivo_actual'), request.files.get('distributivo_anterior')
    if not fa or not fb: return jsonify({'error':'Se requieren ambos archivos'}),400
    pa = os.path.join(app.config['UPLOAD_FOLDER'],'dist_actual'+os.path.splitext(fa.filename)[1])
    pb = os.path.join(app.config['UPLOAD_FOLDER'],'dist_anterior'+os.path.splitext(fb.filename)[1])
    fa.save(pa); fb.save(pb); session['pa']=pa; session['pb']=pb
    try:
        da,db = read_distributivo(pa), read_distributivo(pb)
        for c in ['ESTADO DEL PUESTO','NÚMERO IDENTIFICACIÓN','NOMBRES','RÉGIMEN LABORAL','MODALIDAD LABORAL']:
            if c not in da.columns or c not in db.columns: return jsonify({'error':f'Columna "{c}" no encontrada'}),400
        return jsonify({'success':True,'actual':{'total':len(da),'ocupados':len(filter_ocupados(da)),'filename':fa.filename},
                        'anterior':{'total':len(db),'ocupados':len(filter_ocupados(db)),'filename':fb.filename}})
    except Exception as e: return jsonify({'error':str(e)}),400

@app.route('/analyze')
def analyze():
    try:
        da,db = read_distributivo(session['pa']), read_distributivo(session['pb'])
        sal,ent,camb = compare_distributivos(db,da); cats = get_distribution(da)
        sl = [{'cedula':clean_id(r['NÚMERO IDENTIFICACIÓN']),'nombres':str(r['NOMBRES']).strip(),'escala':str(r.get('ESCALA OCUPACIONAL','')).strip(),'rmu':fmt_rmu(r.get('RMU PUESTO','')),'regimen':str(r.get('RÉGIMEN LABORAL','')).strip(),'modalidad':str(r.get('MODALIDAD LABORAL','')).strip(),'motivo':'RENUNCIA'} for _,r in sal.iterrows()]
        el = [{'cedula':clean_id(r['NÚMERO IDENTIFICACIÓN']),'nombres':str(r['NOMBRES']).strip(),'escala':str(r.get('ESCALA OCUPACIONAL','')).strip(),'rmu':fmt_rmu(r.get('RMU PUESTO','')),'regimen':str(r.get('RÉGIMEN LABORAL','')).strip(),'modalidad':str(r.get('MODALIDAD LABORAL','')).strip()} for _,r in ent.iterrows()]
        cl = [{'cedula':clean_id(r['NÚMERO IDENTIFICACIÓN']),'nombres':str(r['NOMBRES']).strip(),'escala':str(r.get('ESCALA OCUPACIONAL','')).strip(),'rmu':fmt_rmu(r.get('RMU PUESTO','')),'regimen':str(r.get('RÉGIMEN LABORAL','')).strip(),'modalidad':str(r.get('MODALIDAD LABORAL','')).strip()} for _,r in camb.iterrows()]
        cs = {k:{'title':v['title'],'subtitle':v['subtitle'],'count':len(v['data']),'data':df_to_rows(v['data'])} for k,v in cats.items()}
        return jsonify({'salieron':sl,'entraron':el,'cambiaron':cl,'categories':cs,'total_ocupados_actual':len(filter_ocupados(da))})
    except Exception as e:
        import traceback; return jsonify({'error':str(e),'trace':traceback.format_exc()}),500

@app.route('/novedades_plantilla')
def get_novedades(): return jsonify(DEFAULT_NOVEDADES)

@app.route('/previous_data/<nid>')
def get_prev(nid): return jsonify(_load_prev().get(nid,{'rows':[]}))

@app.route('/generate', methods=['POST'])
def generate():
    try:
        data = request.json; dg = data.get('datos_generales',{}); novs = data.get('novedades_manuales',[])
        motivos = data.get('motivos_override',{})
        da,db = read_distributivo(session['pa']), read_distributivo(session['pb'])
        sal,ent,camb = compare_distributivos(db,da); cats = get_distribution(da)
        for ced,mot in motivos.items(): sal.loc[sal['NÚMERO IDENTIFICACIÓN']==ced,'MOTIVO'] = mot
        # N1: personas desvinculadas agregadas manualmente. Se fueron a mitad de mes
        # pero SIGUEN en el distributivo actual, por lo que compare_distributivos no
        # las detecta. Se buscan por cédula en el distributivo actual y se añaden a sal.
        manuales = data.get('desvinculaciones_manuales', [])
        if manuales:
            act_full = filter_ocupados(da)
            extra = []
            for m in manuales:
                ced_m = clean_id(m.get('cedula',''))
                if not ced_m: continue
                match = act_full[act_full['NÚMERO IDENTIFICACIÓN'].apply(lambda x: clean_id(x)==ced_m)]
                if len(match):
                    rr = match.iloc[[0]].copy()
                    rr['MOTIVO'] = m.get('motivo','RENUNCIA')
                    extra.append(rr)
            if extra:
                sal = pd.concat([sal]+extra, ignore_index=True)
        mm = {'renuncia':'RENUNCIA','fallecimiento':'FALLECIMIENTO','notificacion':'NOTIFICACIÓN'}
        JUB_TIPOS = ['JUBILACIÓN OBLIGATORIA','JUBILACIÓN VOLUNTARIA','JUBILACIÓN POR DISCAPACIDAD','JUBILACIÓN POR INCAPACIDAD']
        for nv in novs:
            if nv.get('id') in mm:
                tgt = mm[nv['id']]; flt = sal[sal['MOTIVO']==tgt]
                if len(flt)>0:
                    nv['active']=True; obs_map = nv.get('auto_observations',{})
                    nv['rows']=[[str(i+1),clean_id(r['NÚMERO IDENTIFICACIÓN']),str(r['NOMBRES']).strip(),str(r.get('ESCALA OCUPACIONAL','')).strip(),fmt_rmu(r.get('RMU PUESTO','')),obs_map.get(clean_id(r['NÚMERO IDENTIFICACIÓN']),'')] for i,(_,r) in enumerate(flt.iterrows())]
        # Novedades automáticas de VINCULACIÓN (nuevos en el mes)
        vinc_novs_auto = []
        if len(ent) > 0:
            obs_map = {}
            for nv in novs:
                if nv.get('id') == 'vinculacion_auto': obs_map = nv.get('auto_observations', {})
            rows_ent = [[str(i+1), clean_id(r['NÚMERO IDENTIFICACIÓN']), str(r['NOMBRES']).strip(),
                         str(r.get('ESCALA OCUPACIONAL','')).strip(), fmt_rmu(r.get('RMU PUESTO','')),
                         obs_map.get(clean_id(r['NÚMERO IDENTIFICACIÓN']), '')]
                        for i, (_, r) in enumerate(ent.iterrows())]
            vinc_novs_auto.append({'id':'vinculacion_auto','title':'VINCULACIONES',
                'description':'','active':True,
                'columns':['NRO.','CÉDULA','APELLIDOS Y NOMBRES','ESCALA','RMU','OBSERVACIÓN'],
                'rows':rows_ent,'has_previous':False,'is_auto':True})
        # Novedades automáticas de VINCULACIÓN GANADOR DE CONCURSO (cambio RMU/modalidad)
        if len(camb) > 0:
            obs_map_c = {}
            for nv in novs:
                if nv.get('id') == 'ganador_concurso_auto': obs_map_c = nv.get('auto_observations', {})
            rows_camb = [[str(i+1), clean_id(r['NÚMERO IDENTIFICACIÓN']), str(r['NOMBRES']).strip(),
                          str(r.get('ESCALA OCUPACIONAL','')).strip(), fmt_rmu(r.get('RMU PUESTO','')),
                          obs_map_c.get(clean_id(r['NÚMERO IDENTIFICACIÓN']), 'GANÓ CONCURSO')]
                         for i, (_, r) in enumerate(camb.iterrows())]
            vinc_novs_auto.append({'id':'ganador_concurso_auto','title':'VINCULACIÓN GANADOR DE CONCURSO',
                'description':'','active':True,
                'columns':['NRO.','CÉDULA','APELLIDOS Y NOMBRES','ESCALA','RMU','OBSERVACIÓN'],
                'rows':rows_camb,'has_previous':False,'is_auto':True})

        # Novedades automáticas de jubilación (una por cada tipo que tenga personas)
        jub_novs_auto = []
        for tipo in JUB_TIPOS:
            flt = sal[sal['MOTIVO']==tipo]
            if len(flt)>0:
                obs_map = {}
                for nv in novs:
                    if nv.get('id')=='jubilacion_auto': obs_map = nv.get('auto_observations',{})
                rows = [[str(i+1),str(r['NOMBRES']).strip(),clean_id(r['NÚMERO IDENTIFICACIÓN']),str(r.get('ESCALA OCUPACIONAL','')).strip(),fmt_rmu(r.get('RMU PUESTO','')),obs_map.get(clean_id(r['NÚMERO IDENTIFICACIÓN']),'')]
                        for i,(_,r) in enumerate(flt.iterrows())]
                jub_novs_auto.append({'id':'jub_auto_'+tipo.lower().replace(' ','_'),'title':tipo,'description':'','active':True,'columns':['No','APELLIDOS Y NOMBRES','C.C. No.','CARGO','RMU','OBSERVACIÓN'],'rows':rows,'has_previous':False,'is_auto':True})
        novs = novs + vinc_novs_auto + jub_novs_auto
        ns = dg.get('numero_corto','2026-095'); dg['numero_informe']=f"MINEDUC-17D08-DDTH-{ns}"
        mes,anio = dg.get('mes','MES').upper(), dg.get('anio',datetime.now().year)
        no = ns.split('-')[-1] if '-' in ns else ns
        wf,ef = f'INFORME_{no}_{mes}_{anio}.docx', f'TABLAS_{no}_{mes}_{anio}.xlsx'
        wp,ep = os.path.join(app.config['OUTPUT_FOLDER'],wf), os.path.join(app.config['OUTPUT_FOLDER'],ef)
        generate_word(dg,novs,cats,wp); generate_excel(cats,novs,ep)
        return jsonify({'success':True,'word_file':wf,'excel_file':ef})
    except Exception as e:
        import traceback; return jsonify({'error':str(e),'trace':traceback.format_exc()}),500

@app.route('/download/<fn>')
def download(fn):
    fp = os.path.join(app.config['OUTPUT_FOLDER'],fn)
    return send_file(fp,as_attachment=True) if os.path.exists(fp) else (jsonify({'error':'No encontrado'}),404)

def open_browser():
    webbrowser.open('http://localhost:5000')

if __name__ == '__main__':
    print("\n"+"="*60+"\n  GENERADOR DE INFORMES - DISTRITO 17D08\n  Abriendo navegador en http://localhost:5000\n"+"="*60+"\n")
    threading.Timer(1.5, open_browser).start()
    app.run(debug=False, host='0.0.0.0', port=5000)
